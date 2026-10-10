"""Real synthetic DOCX/OPC archives with explicit source-text/structure oracles."""

import hashlib
import io
import json
from dataclasses import replace
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

from services.ingestion.errors import IngestionError
from services.ingestion.office.docx import parse_docx
from services.ingestion.office.models import OfficeLimits
from services.ingestion.office.opc import OfficePackage

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
STRICT_W = "http://purl.oclc.org/ooxml/wordprocessingml/main"
STRICT_R = "http://purl.oclc.org/ooxml/officeDocument/relationships"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
TYPES = "http://schemas.openxmlformats.org/package/2006/content-types"
MAIN_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"


def paragraph(text: str, properties: str = "") -> str:
    return f'<w:p>{properties}<w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:p>'


def package_file(tmp_path: Path, body: str, *, parts=None, relations=None, strict=False) -> Path:
    """Generate a real bounded container; retain text generators rather than blobs."""
    path = tmp_path / "oracle.docx"
    namespaces = {"w": STRICT_W if strict else W, "r": STRICT_R if strict else R}
    added = dict(parts or {})
    document = '<w:document xmlns:w="{w}" xmlns:r="{r}"><w:body>{body}</w:body></w:document>'.format(**namespaces, body=body)
    for name, value in list(added.items()):
        if isinstance(value, str):
            added[name] = value.replace("{W}", namespaces["w"]).replace("{R}", namespaces["r"])
    records = list(relations or [])
    known_types = {
        "styles.xml": "styles", "numbering.xml": "numbering", "footnotes.xml": "footnotes",
        "endnotes.xml": "endnotes", "comments.xml": "comments", "header1.xml": "header", "footer1.xml": "footer",
    }
    overrides = [f'<Override PartName="/word/document.xml" ContentType="{MAIN_MIME}"/>']
    for name in added:
        basename = name.rsplit("/", 1)[-1]
        if basename in known_types:
            role = known_types[basename]
            content_type = f"application/vnd.openxmlformats-officedocument.wordprocessingml.{role}+xml"
            overrides.append(f'<Override PartName="/{name}" ContentType="{content_type}"/>')
            if not any(record[2] == basename for record in records):
                records.append((role, namespaces["r"] + "/" + role, basename, False))
        elif name == "docProps/core.xml":
            overrides.append(f'<Override PartName="/{name}" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>')
    content = f'<Types xmlns="{TYPES}"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Default Extension="png" ContentType="image/png"/>{"".join(overrides)}</Types>'
    root_rels = f'<Relationships xmlns="{REL}"><Relationship Id="office" Type="{namespaces["r"]}/officeDocument" Target="word/document.xml"/>'
    if "docProps/core.xml" in added:
        root_rels += '<Relationship Id="core" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
    root_rels += "</Relationships>"
    word_rels = f'<Relationships xmlns="{REL}">'
    for identity, kind, target, external in records:
        mode = ' TargetMode="External"' if external else ""
        word_rels += f'<Relationship Id="{identity}" Type="{kind}" Target="{target}"{mode}/>'
    word_rels += "</Relationships>"
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, value in {"[Content_Types].xml": content, "_rels/.rels": root_rels,
                            "word/document.xml": document, "word/_rels/document.xml.rels": word_rels, **added}.items():
            archive.writestr(name, value)
    return path


def read(path, **kwargs):
    with OfficePackage(path, expected_format="docx", limits=kwargs.pop("limits", None)) as package:
        return parse_docx(package, **kwargs)


@pytest.mark.parametrize("strict", [False, True])
def test_source_order_literal_text_repetition_and_stable_locators(tmp_path, strict):
    body = paragraph(" exact é𝄞 ") + '<w:p><w:r><w:t>left</w:t><w:tab/><w:t>right</w:t><w:br/><w:t>last</w:t><w:noBreakHyphen/><w:softHyphen/></w:r></w:p>' + paragraph(" exact é𝄞 ")
    path = package_file(tmp_path, body, strict=strict)
    original = path.read_bytes()
    extracted = read(path)
    blocks = extracted["units"][0]["blocks"]
    assert [block["text"] for block in blocks] == [" exact é𝄞 ", "left\tright\nlast\u2011\u00ad", " exact é𝄞 "]
    assert blocks[0]["id"] != blocks[2]["id"]
    assert [block["locator"]["element_path"] for block in blocks] == [f"/document[1]/body[1]/p[{at}]" for at in (1, 2, 3)]
    assert read(path) == extracted
    assert path.read_bytes() == original
    assert extracted["metadata"]["ooxml_family"] == ("strict" if strict else "transitional")
    assert extracted["status"] == "ready"
    assert "page" not in repr([block["locator"] for block in blocks])


def test_style_inheritance_sections_outline_override_and_cycle(tmp_path):
    styles = '<w:styles xmlns:w="{W}"><w:style w:type="paragraph" w:styleId="Base"><w:name w:val="Titre base"/><w:pPr><w:outlineLvl w:val="0"/></w:pPr></w:style><w:style w:type="paragraph" w:styleId="Derived"><w:basedOn w:val="Base"/></w:style><w:style w:styleId="A"><w:basedOn w:val="B"/></w:style><w:style w:styleId="B"><w:basedOn w:val="A"/></w:style></w:styles>'
    body = paragraph("Chapitre", '<w:pPr><w:pStyle w:val="Derived"/></w:pPr>') + paragraph("texte") + paragraph("Sous-section", '<w:pPr><w:pStyle w:val="Derived"/><w:outlineLvl w:val="1"/></w:pPr>') + paragraph("corps", '<w:pPr><w:pStyle w:val="Derived"/><w:outlineLvl w:val="9"/></w:pPr>') + paragraph("cycle", '<w:pPr><w:pStyle w:val="A"/></w:pPr>')
    extracted = read(package_file(tmp_path, body, parts={"word/styles.xml": styles}))
    blocks = extracted["units"][0]["blocks"]
    sections = extracted["sections"]
    assert [block["kind"] for block in blocks] == ["heading", "paragraph", "heading", "paragraph", "paragraph"]
    assert [section["level"] for section in sections] == [1, 2]
    assert sections[1]["parent_id"] == sections[0]["id"]
    assert sections[0]["block_ids"] == [block["id"] for block in blocks]
    assert blocks[0]["structure"]["style_chain"] == ["Derived", "Base"]
    assert any(warning["code"] == "DOCX_STYLE_CYCLE" for warning in extracted["warnings"])
    assert extracted["status"] == "ready_partial"


def test_numbering_levels_overrides_restart_and_labels_are_not_source_text(tmp_path):
    numbering = '<w:numbering xmlns:w="{W}"><w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl><w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%1.%2)"/></w:lvl></w:abstractNum><w:num w:numId="1"><w:abstractNumId w:val="0"/><w:lvlOverride w:ilvl="0"><w:startOverride w:val="3"/></w:lvlOverride></w:num></w:numbering>'
    levels = (0, 1, 1, 0, 1)
    body = "".join(paragraph("item", f'<w:pPr><w:numPr><w:ilvl w:val="{level}"/><w:numId w:val="1"/></w:numPr></w:pPr>') for level in levels)
    extracted = read(package_file(tmp_path, body, parts={"word/numbering.xml": numbering}))
    blocks = extracted["units"][0]["blocks"]
    assert [block["structure"]["numbering"]["label"] for block in blocks] == ["3.", "3.a)", "3.b)", "4.", "4.a)"]
    assert all(block["text"] == "item" for block in blocks)
    assert all(block["structure"]["numbering"]["label_is_derived"] for block in blocks)
    assert extracted["status"] == "ready"


@pytest.mark.parametrize("strict", [False, True])
def test_final_revision_view_fields_links_and_annotations(tmp_path, strict):
    body = '<w:p><w:r><w:t>A</w:t></w:r><w:del w:id="1" w:author="author"><w:r><w:delText>OLD</w:delText></w:r></w:del><w:ins w:id="2"><w:r><w:t>NEW</w:t></w:r></w:ins><w:r><w:fldChar w:fldCharType="begin"/><w:instrText> DATE </w:instrText><w:fldChar w:fldCharType="separate"/><w:t>cached date</w:t><w:fldChar w:fldCharType="end"/></w:r><w:hyperlink r:id="link"><w:r><w:t>link text</w:t></w:r></w:hyperlink><w:r><w:footnoteReference w:id="7"/></w:r><w:commentRangeStart w:id="9"/><w:r><w:t>Z</w:t></w:r><w:commentRangeEnd w:id="9"/><w:r><w:commentReference w:id="9"/></w:r></w:p><w:del w:id="10">' + paragraph("deleted paragraph") + '</w:del><w:ins w:id="11">' + paragraph("inserted paragraph") + "</w:ins>"
    relations = [("link", (STRICT_R if strict else R) + "/hyperlink", "https://example.invalid/private", True)]
    parts = {"word/footnotes.xml": '<w:footnotes xmlns:w="{W}"><w:footnote w:id="7">' + paragraph("footnote body") + '</w:footnote><w:footnote w:id="8">' + paragraph("second note") + '</w:footnote></w:footnotes>',
             "word/comments.xml": '<w:comments xmlns:w="{W}"><w:comment w:id="9" w:author="Reviewer">' + paragraph("comment body") + '</w:comment></w:comments>'}
    extracted = read(package_file(tmp_path, body, parts=parts, relations=relations, strict=strict))
    main = extracted["units"][0]["blocks"]
    assert [block["text"] for block in main] == ["ANEWcached datelink textZ", "inserted paragraph"]
    assert main[0]["structure"]["revisions"][0]["text"] == "OLD"
    assert main[0]["structure"]["revisions"][0]["included"] is False
    assert [field["instruction"] for field in main[0]["structure"]["fields"] if field["kind"] == "instruction"] == [" DATE "]
    assert main[0]["structure"]["links"][0]["followed"] is False
    assert main[0]["structure"]["links"][0]["relation"]["external"] is True
    assert extracted["metadata"]["revisions"][0]["text"] == "deleted paragraph"
    assert main[1]["locator"]["element_path"] == "/document[1]/body[1]/ins[3]/p[1]"
    notes = [unit for unit in extracted["units"] if unit["metadata"]["role"] == "footnotes"]
    assert [unit["metadata"]["note"]["id"] for unit in notes] == ["7", "8"]
    assert notes[0]["blocks"][0]["id"] != notes[1]["blocks"][0]["id"]
    assert notes[1]["blocks"][0]["locator"]["element_path"] == "/footnotes[1]/footnote[2]/p[1]"
    assert main[0]["structure"]["references"][0]["target_part"] == "word/footnotes.xml"
    assert main[0]["structure"]["references"][0]["target_unit_id"] == notes[0]["id"]


def test_tables_singleton_nested_grid_fusions_and_exact_cells(tmp_path):
    nested = '<w:tbl><w:tblGrid><w:gridCol w:w="500"/></w:tblGrid><w:tr><w:tc>' + paragraph("nested") + "</w:tc></w:tr></w:tbl>"
    table = '<w:tbl><w:tblGrid><w:gridCol w:w="1000"/><w:gridCol w:w="1000"/><w:gridCol w:w="1000"/></w:tblGrid><w:tr><w:trPr><w:tblHeader/></w:trPr><w:tc><w:tcPr><w:gridSpan w:val="2"/><w:vMerge w:val="restart"/></w:tcPr>' + paragraph("Merged") + '</w:tc><w:tc>' + paragraph("right") + '</w:tc></w:tr><w:tr><w:tc><w:tcPr><w:gridSpan w:val="2"/><w:vMerge/></w:tcPr>' + paragraph("") + '</w:tc><w:tc>' + paragraph("before") + nested + paragraph("after") + '</w:tc></w:tr></w:tbl>'
    extracted = read(package_file(tmp_path, paragraph("lead") + table + paragraph("tail")))
    blocks = extracted["units"][0]["blocks"]
    assert [block["kind"] for block in blocks] == ["paragraph", "table", "paragraph"]
    assert blocks[1]["text"] == "Merged\tright\n\tbefore\nnested\nafter"
    structure = blocks[1]["structure"]
    first = structure["rows"][0]["cells"][0]
    continued = structure["rows"][1]["cells"][0]
    assert first["column_span"] == 2 and first["row_span"] == 2
    assert continued["merge_origin_id"] == first["id"] and continued["text"] == ""
    assert structure["rows"][0]["is_header"] is True
    assert structure["rows"][1]["is_header"] is False
    assert structure["column_count"] == 3
    assert len(extracted["tables"]) == 2
    assert len(structure["rows"][1]["cells"][1]["contents"]) == 3
    for binding in blocks[1]["bindings"]:
        cell = next(cell for row in structure["rows"] for cell in row["cells"] if cell["id"] == binding["cell_id"])
        assert blocks[1]["text"][binding["start"]:binding["end"]] == cell["text"]
    assert blocks[1]["text"].count("Merged") == 1
    assert extracted["status"] == "ready"


def test_images_headers_metadata_and_no_caption_inference(tmp_path):
    from PIL import Image

    stream = io.BytesIO()
    Image.new("RGB", (2, 1), (80, 120, 160)).save(stream, "PNG")
    image = stream.getvalue()
    drawing = '<w:p><w:r><w:drawing><wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"><wp:extent cx="914400" cy="457200"/><wp:docPr id="1" name="Diagram" descr="Alt text"/><a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:graphicData><a:blip r:embed="image"/></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>'
    parts = {"word/media/image1.png": image, "word/header1.xml": '<w:hdr xmlns:w="{W}">' + paragraph("header") + "</w:hdr>",
             "docProps/core.xml": '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Source title</dc:title></cp:coreProperties>'}
    rels = [("image", R + "/image", "media/image1.png", False), ("h1", R + "/header", "header1.xml", False), ("h2", R + "/header", "header1.xml", False)]
    body = drawing + paragraph("Figure 1: separate caption") + '<w:sectPr><w:headerReference w:type="default" r:id="h1"/><w:headerReference w:type="even" r:id="h2"/></w:sectPr>'
    extracted = read(package_file(tmp_path, body, parts=parts, relations=rels))
    image_metadata = extracted["units"][0]["blocks"][0]["structure"]["images"][0]
    assert image_metadata["sha256"] == hashlib.sha256(image).hexdigest()
    assert image_metadata["alt_text"] == "Alt text"
    assert image_metadata["extent_emu"] == {"cx": "914400", "cy": "457200"}
    assert "caption" not in image_metadata
    assert len([unit for unit in extracted["units"] if unit["part"] == "word/header1.xml"]) == 1
    assert extracted["metadata"]["core-properties"][0]["value"] == "Source title"
    assert extracted["metadata"]["sections"][0]["element_path"] == "/document[1]/body[1]/sectPr[3]"


def test_unsupported_elements_are_partial_and_external_images_are_not_fetched(tmp_path):
    body = '<w:p><w:r><w:sym w:font="Wingdings" w:char="F041"/></w:r><w:object/><w:r><w:drawing><a:blip xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" r:link="outside"/></w:drawing></w:r></w:p><w:altChunk r:id="html"/>'
    extracted = read(package_file(tmp_path, body, relations=[("outside", R + "/image", "https://example.invalid/image.png", True)]))
    codes = {warning["code"] for warning in extracted["warnings"]}
    assert {"DOCX_SYMBOL_UNMAPPED", "DOCX_ACTIVE_CONTENT_UNSUPPORTED", "DOCX_EXTERNAL_IMAGE_NOT_LOADED", "DOCX_BLOCK_UNSUPPORTED"} <= codes
    assert extracted["status"] == "ready_partial" and extracted["coverage"]["unsupported"] == 4
    assert "sha256" not in extracted["units"][0]["blocks"][0]["structure"]["images"][0]


@pytest.mark.parametrize("limit,value,body", [
    ("max_blocks", 1, paragraph("first") + paragraph("second")),
    ("max_blocks", 1, '<w:tbl><w:tr><w:tc>' + paragraph("cell") + '</w:tc></w:tr></w:tbl>'),
    ("max_cells", 1, '<w:tbl><w:tr><w:tc>' + paragraph("a") + '</w:tc><w:tc>' + paragraph("b") + '</w:tc></w:tr></w:tbl>'),
    ("max_text_chars", 10, '<w:del>' + paragraph("hidden revision too long") + '</w:del>'),
    ("max_cell_chars", 2, '<w:tbl><w:tr><w:tc>' + paragraph("long cell") + '</w:tc></w:tr></w:tbl>'),
])
def test_global_quotas_include_nested_blocks_and_deleted_text(tmp_path, limit, value, body):
    limits = replace(OfficeLimits(), **{limit: value})
    path = package_file(tmp_path, body)
    with pytest.raises(IngestionError) as caught:
        read(path, limits=limits)
    assert caught.value.code == "OFFICE_LIMIT_EXCEEDED"


def test_cooperative_checkpoint_stops_at_exact_unit_and_replay_is_deterministic(tmp_path):
    path = package_file(tmp_path, "".join(paragraph(f"line {index}") for index in range(1000)))
    calls = []

    def pause(key):
        calls.append(key)
        if len(calls) == 7:
            raise IngestionError("CHECKPOINT_REQUESTED", "Pause")

    with pytest.raises(IngestionError) as caught:
        read(path, checkpoint=pause)
    assert caught.value.code == "CHECKPOINT_REQUESTED"
    assert calls[-1] == "word/document.xml:/document[1]/body[1]/p[7]"
    first = read(path)
    assert len(first["units"][0]["blocks"]) == 1000
    assert first == read(path)


def test_style_and_wrapper_ids_do_not_depend_on_repeated_text(tmp_path):
    body = '<w:sdt><w:sdtPr/><w:sdtContent>' + paragraph("same") + '</w:sdtContent></w:sdt><w:customXml>' + paragraph("same") + "</w:customXml>"
    blocks = read(package_file(tmp_path, body))["units"][0]["blocks"]
    assert [block["locator"]["element_path"] for block in blocks] == ["/document[1]/body[1]/sdt[1]/sdtContent[2]/p[1]", "/document[1]/body[1]/customXml[2]/p[1]"]
    assert blocks[0]["id"] != blocks[1]["id"]


def test_actual_python_docx_document_with_picture_table_and_default_properties(tmp_path):
    from docx import Document
    from PIL import Image

    document = Document()
    document.core_properties.title = "Fixture native python-docx"
    document.add_heading("Real generated heading", level=1)
    document.add_paragraph("source\ttab\nnewline")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "singleton"
    picture = io.BytesIO()
    Image.new("RGB", (3, 2), "blue").save(picture, "PNG")
    document.add_picture(picture)
    document.sections[0].header.paragraphs[0].text = "Header source"
    path = tmp_path / "library-generated.docx"
    document.save(path)
    extracted = read(path)
    blocks = extracted["units"][0]["blocks"]
    assert [block["kind"] for block in blocks] == ["heading", "paragraph", "table", "paragraph"]
    assert blocks[1]["text"] == "source\ttab\nnewline"
    assert blocks[2]["structure"]["rows"][0]["cells"][0]["text"] == "singleton"
    assert blocks[3]["structure"]["images"][0]["size_bytes"] > 0
    assert next(value["value"] for value in extracted["metadata"]["core-properties"] if value["name"] == "title") == "Fixture native python-docx"
    assert extracted["sections"][0]["title"] == "Real generated heading"
    assert extracted["status"] == "ready"


def test_table_deleted_rows_header_false_and_legacy_horizontal_merge(tmp_path):
    body = '<w:tbl><w:tr><w:trPr><w:del w:id="3"/></w:trPr><w:tc>' + paragraph("deleted row") + '</w:tc></w:tr><w:tr><w:trPr><w:tblHeader w:val="false"/></w:trPr><w:tc><w:tcPr><w:hMerge w:val="restart"/></w:tcPr>' + paragraph("horizontal origin") + '</w:tc><w:tc><w:tcPr><w:hMerge/></w:tcPr>' + paragraph("") + '</w:tc></w:tr></w:tbl>'
    extracted = read(package_file(tmp_path, body))
    table = extracted["units"][0]["blocks"][0]
    assert table["text"] == "horizontal origin\t"
    assert table["structure"]["rows"][0]["is_header"] is False
    cells = table["structure"]["rows"][0]["cells"]
    assert cells[0]["column_span"] == 2 and cells[1]["merge_origin_id"] == cells[0]["id"]
    assert table["structure"]["column_count"] == 2
    assert extracted["metadata"]["revisions"][0]["text"] == "deleted row"
    assert cells[0]["element_path"] == "/document[1]/body[1]/tbl[1]/tr[2]/tc[2]"


def test_textbox_preserves_final_revision_view_and_exact_source_anchor(tmp_path):
    textbox = '<w:p><w:r><w:t>before</w:t><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml"><v:textbox><w:txbxContent><w:p><w:r><w:t>first</w:t></w:r><w:del><w:r><w:delText>old</w:delText></w:r></w:del></w:p>' + paragraph("second") + '</w:txbxContent></v:textbox></v:shape></w:pict><w:t>after</w:t></w:r></w:p>'
    extracted = read(package_file(tmp_path, textbox))
    blocks = extracted["units"][0]["blocks"]
    assert [block["text"] for block in blocks] == ["beforeafter", "first\nsecond"]
    assert blocks[1]["kind"] == "textbox"
    assert blocks[1]["locator"]["element_path"] == "/document[1]/body[1]/p[1]/r[1]/pict[2]/shape[1]/textbox[1]/txbxContent[1]"
    assert blocks[1]["structure"]["contents"][0]["structure"]["revisions"][0]["text"] == "old"
    assert extracted["status"] == "ready_partial"


def test_duplicate_note_ids_have_distinct_units_and_visible_ambiguity(tmp_path):
    parts = {"word/footnotes.xml": '<w:footnotes xmlns:w="{W}"><w:footnote w:id="1">' + paragraph("first") + '</w:footnote><w:footnote w:id="1">' + paragraph("second") + '</w:footnote></w:footnotes>'}
    extracted = read(package_file(tmp_path, paragraph("main"), parts=parts))
    units = extracted["units"][1:]
    assert len({unit["id"] for unit in units}) == 2
    assert "DOCX_NOTE_ID_DUPLICATE" in {warning["code"] for warning in extracted["warnings"]}
    assert extracted["status"] == "ready_partial"
    with pytest.raises(IngestionError) as caught:
        read(package_file(tmp_path, paragraph("main"), parts=parts), limits=replace(OfficeLimits(), max_units=2))
    assert caught.value.code == "OFFICE_LIMIT_EXCEEDED"


def test_alternate_content_selects_fallback_once_and_keeps_source_path(tmp_path):
    body = '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><mc:Choice Requires="unknown">' + paragraph("choice") + '</mc:Choice><mc:Fallback>' + paragraph("fallback") + '</mc:Fallback></mc:AlternateContent>'
    extracted = read(package_file(tmp_path, body))
    blocks = extracted["units"][0]["blocks"]
    assert [block["text"] for block in blocks] == ["fallback"]
    assert blocks[0]["locator"]["element_path"] == "/document[1]/body[1]/AlternateContent[1]/Fallback[2]/p[1]"
    assert extracted["status"] == "ready_partial"


def test_unknown_inline_word_element_is_not_silently_reported_complete(tmp_path):
    body = '<w:p><w:futureFeature><w:r><w:t>retained text</w:t></w:r></w:futureFeature></w:p>'
    extracted = read(package_file(tmp_path, body))
    assert extracted["units"][0]["blocks"][0]["text"] == "retained text"
    assert extracted["status"] == "ready_partial"
    assert extracted["warnings"][0]["code"] == "DOCX_INLINE_ELEMENT_UNSUPPORTED"


def test_completed_part_cache_restores_sections_tables_revisions_and_numbering(tmp_path):
    numbering = '<w:numbering xmlns:w="{W}"><w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl></w:abstractNum><w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num></w:numbering>'
    properties = '<w:pPr><w:numPr><w:numId w:val="1"/></w:numPr></w:pPr>'
    body = paragraph("section", '<w:pPr><w:outlineLvl w:val="0"/></w:pPr>') + '<w:sdt><w:sdtPr><w:alias w:val="Structured content"/></w:sdtPr><w:sdtContent>' + paragraph("numbered main", properties) + '</w:sdtContent></w:sdt><w:del>' + paragraph("deleted source") + '</w:del><w:tbl><w:tr><w:tc>' + paragraph("cell") + '</w:tc></w:tr></w:tbl>'
    parts = {"word/numbering.xml": numbering,
             "word/header1.xml": '<w:hdr xmlns:w="{W}">' + paragraph("numbered header", properties) + '</w:hdr>',
             "word/footnotes.xml": '<w:footnotes xmlns:w="{W}"><w:footnote w:id="2">' + paragraph("numbered note", properties) + '</w:footnote></w:footnotes>'}
    path = package_file(tmp_path, body, parts=parts)
    caches = {}

    def persist(unit):
        # Exercise the JSON checkpoint boundary, not only shared Python objects.
        caches[unit["id"]] = json.loads(json.dumps(unit))

    fresh = read(path, unit_sink=persist)
    assert len(caches) == 3
    assert all("_resume" not in unit for unit in fresh["units"])
    assert read(path, cached_units=caches) == fresh
    # A completed main part may be restored while later parts are extracted afresh.
    main_cache = {fresh["units"][0]["id"]: caches[fresh["units"][0]["id"]]}
    mixed = read(path, cached_units=main_cache)
    assert mixed == fresh
    numbered = [block["structure"]["numbering"]["label"] for unit in mixed["units"] for block in unit["blocks"] if "numbering" in block["structure"]]
    assert numbered == ["1.", "2.", "3."]
    assert mixed["metadata"]["content_wrappers"] == fresh["metadata"]["content_wrappers"]


def test_checkpoint_written_before_cooperative_pause_on_completed_part(tmp_path):
    parts = {"word/header1.xml": '<w:hdr xmlns:w="{W}">' + paragraph("header") + '</w:hdr>'}
    path = package_file(tmp_path, paragraph("main"), parts=parts)
    caches = {}

    def persist(unit):
        caches[unit["id"]] = json.loads(json.dumps(unit))

    def stop(key):
        if key == "word/document.xml:/document[1]:complete":
            assert len(caches) == 1
            raise IngestionError("CHECKPOINT_REQUESTED", "Pause")

    with pytest.raises(IngestionError) as caught:
        read(path, unit_sink=persist, checkpoint=stop)
    assert caught.value.code == "CHECKPOINT_REQUESTED"
    assert read(path, cached_units=caches) == read(path)


def test_incomplete_or_foreign_checkpoint_is_rejected(tmp_path):
    path = package_file(tmp_path, paragraph("first source"))
    caches = {}
    read(path, unit_sink=lambda unit: caches.update({unit["id"]: json.loads(json.dumps(unit))}))
    incomplete = {identity: {key: value for key, value in unit.items() if key != "_resume"} for identity, unit in caches.items()}
    with pytest.raises(IngestionError) as caught:
        read(path, cached_units=incomplete)
    assert caught.value.code == "OFFICE_INVALID_CHECKPOINT"
    path = package_file(tmp_path, paragraph("second source same ordinal"))
    with pytest.raises(IngestionError) as caught:
        read(path, cached_units=caches)
    assert caught.value.code == "OFFICE_INVALID_CHECKPOINT"


def test_missing_note_target_is_visible_and_cache_does_not_hide_it(tmp_path):
    body = '<w:p><w:r><w:t>main</w:t><w:footnoteReference w:id="404"/></w:r></w:p>'
    path = package_file(tmp_path, body)
    caches = {}
    fresh = read(path, unit_sink=lambda unit: caches.update({unit["id"]: json.loads(json.dumps(unit))}))
    reference = fresh["units"][0]["blocks"][0]["structure"]["references"][0]
    assert reference["target_unit_id"] is None
    assert fresh["status"] == "ready_partial"
    assert "DOCX_REFERENCE_UNRESOLVED" in {warning["code"] for warning in fresh["warnings"]}
    assert read(path, cached_units=caches) == fresh


def test_xml_comments_are_inert_without_changing_raw_sibling_anchor_counts(tmp_path):
    body = '<!-- source comment -->' + paragraph("first") + '<?ignored processing?>' + paragraph("second")
    extracted = read(package_file(tmp_path, body))
    blocks = extracted["units"][0]["blocks"]
    assert [block["text"] for block in blocks] == ["first", "second"]
    assert [block["locator"]["element_path"] for block in blocks] == ["/document[1]/body[1]/p[2]", "/document[1]/body[1]/p[4]"]
    assert extracted["status"] == "ready"


def test_merged_grid_work_is_bounded_before_large_vertical_maps(tmp_path):
    body = '<w:tbl><w:tr><w:tc><w:tcPr><w:gridSpan w:val="1000"/><w:vMerge w:val="restart"/></w:tcPr>' + paragraph("source") + '</w:tc></w:tr><w:tr><w:tc><w:tcPr><w:gridSpan w:val="1000"/><w:vMerge/></w:tcPr>' + paragraph("") + '</w:tc></w:tr></w:tbl>'
    with pytest.raises(IngestionError) as caught:
        read(package_file(tmp_path, body), limits=replace(OfficeLimits(), max_cells=1000))
    assert caught.value.code == "OFFICE_LIMIT_EXCEEDED"


def test_bad_span_and_merge_mode_are_visible_without_invented_merge(tmp_path):
    body = '<w:tbl><w:tr><w:tc><w:tcPr><w:gridSpan w:val="0"/><w:vMerge w:val="unknown"/></w:tcPr>' + paragraph("retained") + '</w:tc></w:tr></w:tbl>'
    extracted = read(package_file(tmp_path, body))
    assert extracted["units"][0]["blocks"][0]["text"] == "retained"
    assert extracted["status"] == "ready_partial"
    assert {"DOCX_TABLE_SPAN_INVALID", "DOCX_MERGE_MODE_INVALID"} <= {warning["code"] for warning in extracted["warnings"]}


def test_paragraph_property_revision_is_retained_as_data(tmp_path):
    body = paragraph("current text", '<w:pPr><w:outlineLvl w:val="0"/><w:pPrChange w:id="77" w:author="old author"><w:pPr><w:outlineLvl w:val="2"/></w:pPr></w:pPrChange></w:pPr>')
    block = read(package_file(tmp_path, body))["units"][0]["blocks"][0]
    assert block["kind"] == "heading" and block["structure"]["outline_level"] == 0
    revision = block["structure"]["property_revisions"][0]
    assert revision["kind"] == "pPrChange"
    assert revision["attributes"][f"{{{W}}}id"] == "77"
    assert revision["children"][0]["children"][0]["attributes"][f"{{{W}}}val"] == "2"
