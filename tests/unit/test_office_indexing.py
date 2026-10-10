"""Real OOXML/SQLite/FTS provenance; embedding, tokenizer and Qdrant are explicit doubles."""

import asyncio
import copy
import hashlib
import json
import runpy
from pathlib import Path
from uuid import UUID, uuid5

import pytest
from docx import Document

from services.api.db import Database, json_dump, now, uid
from services.api.errors import ApiError
from services.api.indexing import CHUNKER_REVISION, Indexer, extraction_content_hash, text_loss
from services.api.settings import Settings
from services.ingestion.office.pipeline import extract_office, office_content_hash

WORKBOOK = runpy.run_path(str(Path(__file__).parents[1] / "fixtures" / "office" / "xlsx_cases.py"))


class ExplicitEmbedding:
    def count(self, text, passage=True):
        return len(text) // 4 + 3

    def embed(self, texts, passage=True):
        return [[1.0] + [0.0] * 383 for _ in texts]


class ExplicitLlmTokenizer:
    def count(self, text):
        return len(text) // 4 + 1


class ExplicitVectors:
    def __init__(self):
        self.points = {}
        self.fail = False

    async def ensure_collection(self):
        pass

    async def upsert(self, points):
        if self.fail:
            raise ApiError("qdrant_unavailable", "Double explicite : Qdrant indisponible.", 503)
        self.points.update({point["id"]: point for point in points})

    async def verify(self, expected):
        assert {key: self.points[key]["payload"]["text_hash"] for key in expected} == expected


@pytest.fixture
def office_storage(tmp_path):
    settings = Settings(tmp_path, {"app": {"data_dir": "runtime", "port": 8785},
                                   "chunking": {"overlap_max_tokens": 0}})
    db = Database(settings.db_path)
    db.initialize()
    vectors = ExplicitVectors()
    return settings, db, vectors, Indexer(db, ExplicitEmbedding(), vectors, settings, ExplicitLlmTokenizer())


def ingest_fixture(storage, document_format, *, text="A😀é ﬁ e\u0301 CCU-21 : tension 72 V.", name="manual"):
    settings, db, _, _ = storage
    originals = settings.data_dir / "originals"
    originals.mkdir(parents=True, exist_ok=True)
    path = originals / (name + "." + document_format)
    if document_format == "docx":
        document = Document()
        document.add_heading("Préambule", level=1)
        document.add_paragraph(text)
        document.add_paragraph(text)
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Valeur structurée 42"
        document.save(path)
    else:
        WORKBOOK["write_workbook"](path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    imported = db.import_original(f"sources/{name}.{document_format}", digest, path)
    extraction = extract_office(path, settings.data_dir / "extractions" / imported["version_id"],
                                settings.profile, imported["version_id"], document_format)
    assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
    return imported, extraction


def new_job(db, imported):
    job_id = uid()
    db.execute("INSERT INTO jobs(id,document_id,version_id,state,stage,created_at,updated_at) VALUES(?,?,?,'queued','queued',?,?)",
               (job_id, imported["document_id"], imported["version_id"], now(), now()))
    return job_id


def reseal(extraction, version_id):
    extraction["source_hash"] = office_content_hash(extraction)
    extraction["extraction_revision_id"] = str(uuid5(UUID(version_id), extraction["pipeline_fingerprint"] + ":" + extraction["source_hash"]))
    for unit in extraction["units"]:
        for block in unit["blocks"]:
            block["extraction_revision_id"] = extraction["extraction_revision_id"]


@pytest.mark.parametrize("document_format", ["docx", "xlsx"])
def test_true_office_stage_is_idempotent_and_has_no_pdf_pages(office_storage, document_format):
    imported, extraction = ingest_fixture(office_storage, document_format)
    _, db, _, indexer = office_storage
    generation, chunks = indexer.stage(imported["job_id"], extraction)
    assert chunks
    before = {table: db.one(f"SELECT count(*) n FROM {table}")["n"]
              for table in ("blocks", "chunks", "office_units", "office_unit_blocks", "office_cells", "office_cell_bindings")}
    assert indexer.stage(imported["job_id"], extraction) == (generation, chunks)
    assert before == {table: db.one(f"SELECT count(*) n FROM {table}")["n"] for table in before}
    assert db.rows("SELECT * FROM pages") == []
    assert db.rows("SELECT * FROM blocks WHERE page_index IS NOT NULL") == []
    assert db.rows("SELECT * FROM chunk_sources WHERE page_index IS NOT NULL") == []
    assert db.version(imported["version_id"])["page_count"] is None
    assert db.rows("PRAGMA foreign_key_check") == []
    assert db.one("PRAGMA integrity_check")["integrity_check"] == "ok"
    original_blocks = {block["id"]: block for unit in extraction["units"] for block in unit["blocks"]}
    for stored in db.rows("SELECT * FROM blocks"):
        original = original_blocks[stored["id"]]
        assert stored["text"] == original["text"]
        assert stored["source_text_hash"] == hashlib.sha256(original["text"].encode()).hexdigest()
        assert stored["extraction_revision_id"] == extraction["extraction_revision_id"]
        metadata = json.loads(stored["metadata_json"])
        assert metadata["locator"] == original["locator"]
        assert metadata["structure"] == original["structure"]
        assert metadata["bindings"] == original["bindings"]
        assert metadata["format"] == document_format
    assert db.one("SELECT source_hash FROM extraction_revisions")["source_hash"] == extraction_content_hash(extraction)


def test_docx_unicode_source_offsets_and_repeated_anchors(office_storage):
    text = "A😀é ﬁ e\u0301 CCU-21 72 V. " * 140
    imported, extraction = ingest_fixture(office_storage, "docx", text=text)
    _, db, _, indexer = office_storage
    _, chunks = indexer.stage(imported["job_id"], extraction)
    paragraphs = [block for block in extraction["units"][0]["blocks"] if block["text"] == text]
    assert len(paragraphs) == 2 and paragraphs[0]["id"] != paragraphs[1]["id"]
    assert paragraphs[0]["locator"]["element_path"] != paragraphs[1]["locator"]["element_path"]
    for chunk in chunks:
        sources = db.rows("SELECT s.*,b.text FROM chunk_sources s JOIN chunks c ON c.chunk_uuid=s.chunk_uuid JOIN blocks b ON b.generation_id=c.generation_id AND b.id=s.block_id WHERE s.chunk_uuid=? ORDER BY s.position", (chunk["id"],))
        assert "\n".join(source["text"][source["start_offset"]:source["end_offset"]] for source in sources) == chunk["text"]
        assert chunk["tokens"] <= 320
        assert chunk["llm_tokens"] == ExplicitLlmTokenizer().count(chunk["text"])
    assert len([chunk for chunk in chunks if chunk["block_id"] == paragraphs[0]["id"]]) > 1


def test_xlsx_sparse_cells_and_native_binding_offsets_are_preserved(office_storage):
    imported, extraction = ingest_fixture(office_storage, "xlsx")
    _, db, _, indexer = office_storage
    _, chunks = indexer.stage(imported["job_id"], extraction)
    assert db.one("SELECT count(*) n FROM office_cells")["n"] == 27
    first = extraction["units"][0]
    native = {cell["address"]: cell for cell in first["cells"]}
    for stored in db.rows("SELECT * FROM office_cells WHERE unit_id=?", (first["id"],)):
        assert json.loads(stored["data_json"]) == native[stored["address"]]
    assert native["B2"]["value"]["value"] == "9007199254740993"
    assert native["C2"]["cached_value"]["value"] == "18014398509481986"
    assert native["C3"]["cached_value"] is None and native["C3"]["cache_freshness"] == "absent"
    assert db.one("SELECT address FROM office_cells WHERE row_index=1048576 AND column_index=16384")["address"] == "XFD1048576"
    bindings = db.rows("SELECT x.*,b.text FROM office_cell_bindings x JOIN blocks b ON b.generation_id=x.generation_id AND b.id=x.block_id")
    expected = [binding for unit in extraction["units"] for block in unit["blocks"] for binding in block["bindings"]]
    assert len(bindings) == len(expected)
    assert all(binding["text"][binding["start_offset"]:binding["end_offset"]] for binding in bindings)
    for chunk in chunks:
        unit_ids = {json.loads(row["metadata_json"])["locator"]["unit_id"] for row in db.rows(
            "SELECT b.metadata_json FROM chunk_sources s JOIN blocks b ON b.generation_id=? AND b.id=s.block_id WHERE s.chunk_uuid=?",
            (db.one("SELECT generation_id FROM jobs WHERE id=?", (imported["job_id"],))["generation_id"], chunk["id"]))}
        assert unit_ids == {chunk["unit_id"]}


@pytest.mark.parametrize("document_format", ["docx", "xlsx"])
def test_office_indexing_uses_empty_vector_page_indices(office_storage, document_format):
    imported, extraction = ingest_fixture(office_storage, document_format)
    _, db, vectors, indexer = office_storage
    generation = asyncio.run(indexer.index(imported["job_id"], extraction))
    assert db.one("SELECT active_generation_id FROM documents")["active_generation_id"] == generation
    assert vectors.points
    assert all(point["payload"]["page_indices"] == [] for point in vectors.points.values())
    assert all(point["payload"]["block_ids"] for point in vectors.points.values())
    if document_format == "docx":
        assert db.one("SELECT count(*) n FROM chunks_fts WHERE chunks_fts MATCH 'CCU'")["n"] > 0


def test_office_partial_publication_requires_explicit_acceptance(office_storage):
    imported, extraction = ingest_fixture(office_storage, "docx")
    _, db, _, indexer = office_storage
    first = asyncio.run(indexer.index(imported["job_id"], extraction))
    second_job = new_job(db, imported)
    partial = copy.deepcopy(extraction)
    partial["status"] = "ready_partial"
    partial["coverage"]["unsupported"] = 1
    assert partial["parser_complete"] is True and text_loss(partial)
    second = asyncio.run(indexer.index(second_job, partial))
    assert db.one("SELECT active_generation_id FROM documents")["active_generation_id"] == first
    assert db.one("SELECT published_at FROM index_generations WHERE id=?", (second,))["published_at"] is None
    count = db.one("SELECT expected_chunks FROM index_generations WHERE id=?", (second,))["expected_chunks"]
    indexer.publish(second_job, second, count, partial=True, allow_partial=True)
    assert db.one("SELECT active_generation_id FROM documents")["active_generation_id"] == second
    assert db.one("SELECT count(*) n FROM blocks WHERE generation_id=?", (first,))["n"] > 0


def test_office_vector_failure_preserves_previous_generation(office_storage):
    imported, extraction = ingest_fixture(office_storage, "xlsx")
    _, db, vectors, indexer = office_storage
    first = asyncio.run(indexer.index(imported["job_id"], extraction))
    second_job = new_job(db, imported)
    vectors.fail = True
    with pytest.raises(ApiError, match="Qdrant indisponible"):
        asyncio.run(indexer.index(second_job, extraction))
    assert db.one("SELECT active_generation_id FROM documents")["active_generation_id"] == first
    assert db.one("SELECT state FROM index_generations WHERE id=(SELECT generation_id FROM jobs WHERE id=?)", (second_job,))["state"] == "staging"


def test_office_retry_reuses_staged_chunk_ids_and_cached_embeddings(office_storage):
    imported, extraction = ingest_fixture(office_storage, "xlsx")
    _, db, vectors, indexer = office_storage
    vectors.fail = True
    with pytest.raises(ApiError):
        asyncio.run(indexer.index(imported["job_id"], extraction))
    generation = db.one("SELECT generation_id FROM jobs WHERE id=?", (imported["job_id"],))["generation_id"]
    staged = db.rows("SELECT chunk_uuid,text_hash FROM chunks WHERE generation_id=? ORDER BY chunk_uuid", (generation,))
    vectors.fail = False
    assert asyncio.run(indexer.index(imported["job_id"], extraction)) == generation
    assert db.rows("SELECT chunk_uuid,text_hash FROM chunks WHERE generation_id=? ORDER BY chunk_uuid", (generation,)) == staged
    assert indexer.diagnostics()["cache_hits"] == len(staged)
    assert db.one("SELECT active_generation_id FROM documents")["active_generation_id"] == generation
    assert db.rows("PRAGMA foreign_key_check") == []


@pytest.mark.parametrize("damage", ["metadata", "coverage", "fingerprint"])
def test_office_staged_generation_refuses_changed_resume_identity(office_storage, damage):
    imported, extraction = ingest_fixture(office_storage, "docx")
    _, db, _, indexer = office_storage
    generation, _ = indexer.stage(imported["job_id"], extraction)
    before = {table: db.rows(f"SELECT * FROM {table}") for table in
              ("index_generations", "extraction_revisions", "blocks", "chunks", "office_documents")}
    changed = copy.deepcopy(extraction)
    if damage == "metadata":
        changed["metadata"]["tampered"] = True
    elif damage == "coverage":
        changed["coverage"]["unsupported"] += 1
    else:
        changed["pipeline_fingerprint"] = "different-parser"
    reseal(changed, imported["version_id"])
    with pytest.raises(ApiError) as caught:
        indexer.stage(imported["job_id"], changed)
    assert caught.value.code == "generation_drift"
    assert {table: db.rows(f"SELECT * FROM {table}") for table in before} == before
    assert db.one("SELECT generation_id FROM jobs WHERE id=?", (imported["job_id"],))["generation_id"] == generation


@pytest.mark.parametrize("damage", ["source_hash", "revision", "locator", "duplicate_cell", "binding"])
def test_office_invalid_sources_are_rejected_before_persistence(office_storage, damage):
    imported, extraction = ingest_fixture(office_storage, "xlsx")
    _, db, _, indexer = office_storage
    if damage == "source_hash":
        extraction["units"][0]["blocks"][0]["source_text_hash"] = "0" * 64
    elif damage == "revision":
        extraction["extraction_revision_id"] = uid()
    elif damage == "locator":
        extraction["units"][0]["blocks"][0]["locator"]["sheet_id"] = "other"
    elif damage == "duplicate_cell":
        extraction["units"][0]["cells"].append(copy.deepcopy(extraction["units"][0]["cells"][0]))
    else:
        extraction["units"][0]["blocks"][0]["bindings"][0]["end"] = 10**6
    if damage != "revision":
        reseal(extraction, imported["version_id"])
    with pytest.raises(ApiError):
        indexer.stage(imported["job_id"], extraction)
    assert db.rows("SELECT * FROM index_generations") == []
    assert db.rows("SELECT * FROM extraction_revisions") == []
    assert db.rows("SELECT * FROM office_cells") == []
    assert db.one("SELECT state FROM jobs WHERE id=?", (imported["job_id"],))["state"] == "queued"


def test_pdf_generation_fingerprint_remains_byte_equivalent(office_storage):
    _, _, _, indexer = office_storage
    expected = hashlib.sha256(json_dump({"extraction": "fixture-pdf-v1",
        "embedding": {"fingerprint": "explicit-test-embedding"}, "llm_tokenizer": "explicit-test-tokenizer",
        "chunking": {"overlap_max_tokens": 0}, "chunker_revision": CHUNKER_REVISION}).encode()).hexdigest()
    assert indexer.generation_fingerprint("fixture-pdf-v1") == expected
    assert indexer.generation_fingerprint("fixture-pdf-v1", "pdf") == expected
    assert len({expected, indexer.generation_fingerprint("fixture-pdf-v1", "docx"),
                indexer.generation_fingerprint("fixture-pdf-v1", "xlsx")}) == 3


def test_revision_snapshot_keeps_metadata_from_archived_generation(office_storage):
    imported, extraction = ingest_fixture(office_storage, "docx")
    _, db, _, indexer = office_storage
    first = asyncio.run(indexer.index(imported["job_id"], extraction))
    previous = db.one("SELECT metadata_json FROM office_documents WHERE generation_id=?", (first,))["metadata_json"]
    changed = copy.deepcopy(extraction)
    changed["metadata"]["new_parser_annotation"] = "new generation"
    reseal(changed, imported["version_id"])
    second = asyncio.run(indexer.index(new_job(db, imported), changed))
    assert first != second
    assert db.one("SELECT metadata_json FROM office_documents WHERE generation_id=?", (first,))["metadata_json"] == previous
    assert "new_parser_annotation" not in json.loads(previous)
    assert "new_parser_annotation" in json.loads(db.one("SELECT metadata_json FROM office_documents WHERE generation_id=?", (second,))["metadata_json"])
