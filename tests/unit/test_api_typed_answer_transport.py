"""Real ASGI/SQLite/OOXML transport; model, E5, vectors and tokenizer are explicit doubles.

The PDF is a real blank container with an explicitly substituted extraction. These tests
exercise publication and replay, not model quality, PDF extraction or a native listener.
"""

import asyncio
import hashlib
import io
import json
import threading
from contextlib import contextmanager
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook

from services.api.main import create_app
from services.api.settings import Settings
from services.ingestion.office.pipeline import extract_office
from tests.integration.test_api_http import (
    CONTROL,
    FakeEmbedding,
    FakeGovernor,
    FakeTokenizer,
    FakeVectors,
    browser_client,
)

INTERNAL_EVENTS = {"structured_answer", "structured_delta"}


def mixed_plan(catalogue):
    """Choose references only; the provider double supplies no documentary values or prose."""
    selected = []
    excerpt_documents = set()
    for reference in catalogue["references"]:
        if reference["kind"] == "excerpt" and reference["document_id"] not in excerpt_documents:
            excerpt_documents.add(reference["document_id"])
            selected.append(reference["id"])
        elif reference["kind"] == "cell" and reference["cell_facts"]["address"] in {"B2", "C2"}:
            selected.append(reference["id"])
    assert len(excerpt_documents) == 2 and len(selected) == 4, "The real retained context must carry all three formats."
    return {"references": selected, "comparisons": []}


class ScriptedProviderDouble:
    """A controllable provider; production QueryService, catalogue validation and renderer stay real."""

    def __init__(self, reply=mixed_plan, *, finish_reason="stop", pause=False, fail=False):
        self.reply, self.finish_reason, self.pause, self.fail = reply, finish_reason, pause, fail
        self.calls = []
        self.raw_by_query = {}
        self.catalogues = {}
        self.paused = threading.Event()
        self.gate = None
        self.db = None

    async def stream(self, messages, cancel_event, output_tokens=None, *, response_format=None):
        self.calls.append({"messages": messages, "output_tokens": output_tokens, "response_format": response_format})
        running = self.db.rows("SELECT id FROM query_runs WHERE state='running'")
        assert len(running) == 1
        query_id = running[0]["id"]
        if response_format is None:
            # A named legacy prose double, used only for pure PDF/DOCX scopes.
            yield {"type": "delta", "text": "La tension nominale est 72 V [S001]."}
            yield {"type": "done", "finish_reason": "stop", "metrics": {"eval_count": 12}}
            return
        catalogue_row = self.db.one("SELECT data_json FROM events WHERE query_id=? AND type='structured_answer'", (query_id,))
        assert catalogue_row is not None, "The complete audit catalogue must precede the provider invocation."
        catalogue = json.loads(catalogue_row["data_json"])["catalogue"]
        self.catalogues[query_id] = catalogue
        reply = self.reply(catalogue) if callable(self.reply) else self.reply
        raw = json.dumps(reply, ensure_ascii=False) if isinstance(reply, dict) else reply
        self.raw_by_query[query_id] = ""
        cut = max(1, len(raw) // 3)
        fragments = [raw[:cut], raw[cut:2 * cut], raw[2 * cut:]]
        for position, fragment in enumerate(fragments):
            self.raw_by_query[query_id] += fragment
            yield {"type": "delta", "text": fragment}
            if self.pause and position == 0:
                self.gate = asyncio.Event()
                self.paused.set()
                await self.gate.wait()
        if self.fail:
            raise RuntimeError("PRIVATE_PROVIDER_ERROR")
        yield {"type": "done", "finish_reason": self.finish_reason,
               "metrics": {"prompt_eval_count": FakeTokenizer().count_messages(messages), "eval_count": len(raw) // 4}}

    async def unblock(self):
        assert self.gate is not None
        self.gate.set()

    async def close(self):
        if self.gate is not None:
            self.gate.set()


@contextmanager
def typed_workspace(tmp_path, monkeypatch, provider, *, literals=False):
    """Only a fresh pytest directory; no worker, resident model or provisioned DATA."""
    monkeypatch.setenv("RAG_CONTROL_TOKEN", CONTROL)
    monkeypatch.setenv("RAG_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("RAG_DB_PATH", str(tmp_path / "data" / "app.sqlite3"))
    settings = Settings(tmp_path, {"app": {"data_dir": "data", "port": 8785},
                                   "chunking": {"overlap_max_tokens": 0}})
    app = create_app(settings=settings, embedding=FakeEmbedding(), vectors=FakeVectors(),
                     tokenizer=FakeTokenizer(), ollama=provider, governor=FakeGovernor(), start_jobs=False)
    provider.db = app.state.db
    with browser_client(app) as client:
        documents = {}
        # A complete blank PDF: catalogue, pages, page, empty stream and byte-accurate xref.
        objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
                   b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R >>",
                   b"<< /Length 0 >>\nstream\n\nendstream"]
        pdf_bytes = io.BytesIO(b"%PDF-1.4\n")
        pdf_bytes.seek(0, 2)
        offsets = []
        for index, data in enumerate(objects, 1):
            offsets.append(pdf_bytes.tell())
            pdf_bytes.write(f"{index} 0 obj\n".encode() + data + b"\nendobj\n")
        xref = pdf_bytes.tell()
        pdf_bytes.write(b"xref\n0 5\n0000000000 65535 f \n")
        for offset in offsets:
            pdf_bytes.write(f"{offset:010} 00000 n \n".encode())
        pdf_bytes.write(f"trailer\n<< /Size 5 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
        document = Document()
        document.add_paragraph("CCU-21 : tension nominale 72 V. A😀é.")
        document.add_paragraph("CCU-21 : la tension du module est 72 V.")
        docx_bytes = io.BytesIO()
        document.save(docx_bytes)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Mesures été"
        sheet.append(["Référence", "Tension nominale", "Longueur"])
        literal = "A😀 & <b> **[S999]** [S001,S002] ![x](https://example.invalid/x)"
        sheet.append(["CCU-21", f"72 V {literal}" if literals else "72 V", "=LEN(A2)"])
        sheet.append(["CCU-22", "OUTSIDE_SCOPE_999V", "=LEN(A3)"])
        xlsx_bytes = io.BytesIO()
        workbook.save(xlsx_bytes)
        workbook.close()
        for document_format, payload in (("pdf", pdf_bytes.getvalue()), ("docx", docx_bytes.getvalue()),
                                         ("xlsx", xlsx_bytes.getvalue())):
            response = client.post("/api/v1/documents/import", files={
                "files": (f"manual.{document_format}", payload, "application/octet-stream")})
            assert response.status_code == 202, response.text
            imported = response.json()
            if literals:
                # A stored synthetic display name; the safe import path itself stays unmodified.
                app.state.db.execute("UPDATE documents SET name=? WHERE id=?",
                                     (f"Source {literal}", imported["document_id"]))
            original = app.state.db.version(imported["version_id"])["blob_path"]
            if document_format == "pdf":
                # Explicit PDF extraction double; the original is an actual blank PDF.
                extraction = {"sha256": hashlib.sha256(payload).hexdigest(), "fingerprint": "pdf-extraction-double",
                              "page_count": 1, "pages": [{"page_index": 0, "width": 595, "height": 842,
                              "blocks": [{"id": "pdf-text", "text": "CCU-21 : tension nominale 72 V.",
                                          "raw_text": "CCU-21 : tension nominale 72 V.", "bbox": [10, 10, 200, 40]}]}]}
            else:
                extraction = extract_office(original, settings.data_dir / "extractions" / imported["version_id"],
                                            settings.profile, imported["version_id"], document_format)
            client.portal.call(app.state.indexer.index, imported["job_id"], extraction)
            documents[document_format] = imported
        yield SimpleNamespace(app=app, client=client, db=app.state.db, documents=documents, provider=provider)


def stored_events(workspace, query_id):
    return [{"id": row["id"], "type": row["type"], "data": json.loads(row["data_json"])}
            for row in workspace.db.rows("SELECT id,type,data_json FROM events WHERE query_id=? ORDER BY id", (query_id,))]


def sse_events(text):
    """Independent parser for the actual HTTP response, including unfamiliar event names."""
    parsed = []
    for frame in text.replace("\r\n", "\n").split("\n\n"):
        fields = {}
        for line in frame.splitlines():
            if line.startswith(("id: ", "event: ", "data: ")):
                key, value = line.split(": ", 1)
                fields[key] = value
        if "data" in fields:
            parsed.append({"id": int(fields["id"]), "type": fields["event"], "data": json.loads(fields["data"])})
    return parsed


def await_query(workspace, query_id):
    async def finish():
        task = workspace.app.state.queries.tasks.get(query_id)
        if task is not None:
            await asyncio.wait_for(asyncio.shield(task), 30)
        await asyncio.sleep(0)
    workspace.client.portal.call(finish)
    assert workspace.db.one("SELECT state FROM query_runs WHERE id=?", (query_id,))["state"] not in {"queued", "running"}


def post_query(workspace, *, scope=None, mode="question"):
    scope = scope or {"kind": "documents", "documentIds": [item["document_id"] for item in workspace.documents.values()]}
    response = workspace.client.post("/api/v1/queries", json={
        "question": "Décrire la tension nominale CCU-21 et la cellule C2 des documents.", "mode": mode, "scope": scope})
    assert response.status_code == 202, response.text
    return response.json()["query_id"]


def database_snapshot(db):
    """All persisted tables, in total column order including WITHOUT ROWID FTS configuration."""
    tables = db.rows("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name")
    result = {}
    for table in tables:
        name = table["name"]
        columns = db.rows(f'PRAGMA table_info("{name}")')
        order = ",".join('"' + column["name"].replace('"', '""') + '"' for column in columns)
        result[name] = db.rows(f'SELECT * FROM "{name}" ORDER BY {order}')
    return result


def assert_public_replay(workspace, query_id):
    stored = stored_events(workspace, query_id)
    expected = [event for event in stored if event["type"] not in INTERNAL_EVENTS]
    response = workspace.client.get(f"/api/v1/queries/{query_id}/events", params={"after": 0})
    assert response.status_code == 200
    assert "event: structured_answer\n" not in response.text and "event: structured_delta\n" not in response.text
    assert sse_events(response.text) == expected
    internal_ids = [event["id"] for event in stored if event["type"] in INTERNAL_EVENTS]
    assert internal_ids, "This oracle must exercise skipped internal events rather than an empty legacy stream."
    cursor = internal_ids[-1]
    expected_tail = [event for event in expected if event["id"] > cursor]
    for params, headers in (({"after": cursor}, {}), ({"after": 0}, {"Last-Event-ID": str(cursor)})):
        tail = workspace.client.get(f"/api/v1/queries/{query_id}/events", params=params, headers=headers)
        assert tail.status_code == 200 and sse_events(tail.text) == expected_tail
    last_id = stored[-1]["id"]
    ended = workspace.client.get(f"/api/v1/queries/{query_id}/events", params={"after": last_id})
    assert ended.status_code == 200 and sse_events(ended.text) == []
    return stored


def assert_raw_is_internal(workspace, query_id):
    stored = assert_public_replay(workspace, query_id)
    raw = "".join(event["data"]["text"] for event in stored if event["type"] == "structured_delta")
    assert raw == workspace.provider.raw_by_query[query_id]
    assert sum(event["type"] == "structured_delta" for event in stored) >= 1
    assert [event["id"] for event in stored] == list(range(1, len(stored) + 1))
    return stored


@pytest.mark.parametrize("mode", ["question", "analysis", "comparison"])
def test_typed_answer_publishes_only_validated_text_with_exact_sql_sse_and_registry(tmp_path, monkeypatch, mode):
    provider = ScriptedProviderDouble()
    with typed_workspace(tmp_path, monkeypatch, provider) as workspace:
        query_id = post_query(workspace, mode=mode)
        await_query(workspace, query_id)
        stored = assert_raw_is_internal(workspace, query_id)
        assert len(provider.calls) == 1 and provider.calls[0]["response_format"] is not None
        evidence = next(event["data"] for event in stored if event["type"] == "structured_answer")
        assert provider.calls[0]["response_format"] == evidence["schema"]
        assert evidence["schema"]["properties"]["references"]["items"]["enum"] == [
            reference["id"] for reference in evidence["catalogue"]["references"]]
        for value, hash_key in ((evidence["schema"], "schema_sha256"), (evidence["catalogue"], "catalogue_sha256")):
            encoded = json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()
            assert hashlib.sha256(encoded).hexdigest() == evidence[hash_key]
        public_sources = next(event["data"]["sources"] for event in stored if event["type"] == "sources")
        source_ids = {source["source_id"] for source in public_sources}
        assert all(set(reference["source_ids"]) <= source_ids for reference in evidence["catalogue"]["references"])
        names = [event["type"] for event in stored]
        assert names.index("sources") < names.index("delta")
        assert names.index("structured_answer") < names.index("structured_delta") < names.index("delta")
        assert names.count("delta") == names.count("done") == 1 and not {"error", "cancelled"}.intersection(names)
        delta = next(event["data"]["text"] for event in stored if event["type"] == "delta")
        done = stored[-1]["data"]
        query = workspace.db.one("SELECT answer,state,metrics_json FROM query_runs WHERE id=?", (query_id,))
        assistant = workspace.db.rows("SELECT content FROM messages WHERE query_id=? AND role='assistant'", (query_id,))
        assert query["state"] == done["status"] == "done"
        assert delta == done["text"] == done["message"] == query["answer"] == assistant[0]["content"]
        assert len(assistant) == 1 and delta != provider.raw_by_query[query_id]
        assert "72 V" in delta and "LEN(A2)" in delta
        assert "OUTSIDE_SCOPE_999V" not in delta
        assert done["structured_answer"]["status"] == "validated"
        assert done["structured_answer"]["rendered_tokens"] == FakeTokenizer().count(delta)
        assert json.loads(query["metrics_json"]) == done["metrics"]
        assert {workspace.db.version(source["version_id"])["format"] for source in done["citations"]} == {"pdf", "docx", "xlsx"}
        for source in done["citations"]:
            registered = workspace.client.get(f"/api/v1/citations/{query_id}/{source['source_id']}")
            assert registered.status_code == 200
            actual = registered.json()
            for key in ("query_id", "source_id", "document_id", "version_id", "generation_id", "extraction_revision_id"):
                assert actual[key] == source[key]
            assert actual.get("locator") == source.get("locator")
            assert actual["version_id"] == workspace.documents[workspace.db.version(source["version_id"])["format"]]["version_id"]
            assert f"[{source['source_id']}]" in delta
        assert workspace.app.state.queries.tasks == {} and workspace.app.state.queries.cancel_events == {}


def assert_no_unvalidated_answer(workspace, query_id, terminal):
    stored = assert_raw_is_internal(workspace, query_id)
    assert stored[-1]["type"] == terminal
    assert all(event["type"] != "delta" for event in stored)
    query = workspace.db.one("SELECT answer FROM query_runs WHERE id=?", (query_id,))
    assert query["answer"] == ""
    assert all(not row["content"] for row in workspace.db.rows(
        "SELECT content FROM messages WHERE query_id=? AND role='assistant'", (query_id,)))
    assert not stored[-1]["data"].get("text", "") and not stored[-1]["data"].get("plan")
    metadata = stored[-1]["data"].get("structured_answer", {})
    assert "plan" not in metadata and "catalogue" not in metadata
    detail = workspace.client.get(f"/api/v1/queries/{query_id}")
    assert detail.status_code == 200 and detail.json()["answer"] == ""
    public = json.dumps(detail.json(), ensure_ascii=False)
    assert "PRIVATE_RAW_MARKER" not in public and "PRIVATE_PROVIDER_ERROR" not in public
    assert workspace.provider.raw_by_query[query_id] not in public
    metrics = detail.json()["metrics"].get("structured_answer", {})
    assert "plan" not in metrics and "catalogue" not in metrics
    assert len(workspace.provider.calls) == 1, "Invalid/truncated/cancelled typed output must not trigger prose fallback or retry."
    return stored


def invalid_reply(kind):
    def reply(catalogue):
        valid = mixed_plan(catalogue)
        if kind == "foreign-reference":
            return {"references": ["R999999"], "comparisons": []}
        if kind == "extra-value":
            return {**valid, "result": 6, "prose": "PRIVATE_RAW_MARKER"}
        if kind == "duplicate-reference":
            return {"references": [valid["references"][0]] * 2, "comparisons": []}
        if kind == "wrong-type":
            return {"references": "PRIVATE_RAW_MARKER", "comparisons": []}
        raise AssertionError(kind)
    return reply


@pytest.mark.parametrize("reply", ["{PRIVATE_RAW_MARKER invalid json", *(invalid_reply(kind) for kind in (
    "foreign-reference", "extra-value", "duplicate-reference", "wrong-type"))],
    ids=["invalid-json", "foreign-reference", "extra-value-and-prose", "duplicate-reference", "wrong-type"])
def test_invalid_typed_output_stays_internal_and_fails_without_fallback(tmp_path, monkeypatch, reply):
    provider = ScriptedProviderDouble(reply)
    with typed_workspace(tmp_path, monkeypatch, provider) as workspace:
        query_id = post_query(workspace)
        await_query(workspace, query_id)
        stored = assert_no_unvalidated_answer(workspace, query_id, "error")
        assert stored[-1]["data"]["code"] == "structured_answer_invalid"
        assert workspace.db.one("SELECT state FROM query_runs WHERE id=?", (query_id,))["state"] == "error"
        assert workspace.db.rows("SELECT * FROM messages WHERE query_id=? AND role='assistant'", (query_id,)) == []


def test_complete_json_at_model_length_limit_is_incomplete_and_never_rendered(tmp_path, monkeypatch):
    provider = ScriptedProviderDouble(finish_reason="length")
    with typed_workspace(tmp_path, monkeypatch, provider) as workspace:
        query_id = post_query(workspace)
        await_query(workspace, query_id)
        stored = assert_no_unvalidated_answer(workspace, query_id, "done")
        assert json.loads(provider.raw_by_query[query_id])["references"], "Exercise a syntactically complete plan at the limit."
        assert stored[-1]["data"]["status"] == "length_limited"
        assert stored[-1]["data"]["finish_reason"] == "length"
        assert stored[-1]["data"]["structured_answer"]["status"] == "incomplete"
        assert workspace.db.one("SELECT state FROM query_runs WHERE id=?", (query_id,))["state"] == "length_limited"


def test_provider_failure_after_raw_fragments_preserves_audit_without_answer(tmp_path, monkeypatch):
    provider = ScriptedProviderDouble(fail=True)
    with typed_workspace(tmp_path, monkeypatch, provider) as workspace:
        query_id = post_query(workspace)
        await_query(workspace, query_id)
        stored = assert_no_unvalidated_answer(workspace, query_id, "error")
        assert sum(event["type"] == "structured_delta" for event in stored) == 3
        assert stored[-1]["data"]["code"] == "query_failed"
        assert "PRIVATE_PROVIDER_ERROR" not in stored[-1]["data"]["message"]


@pytest.mark.parametrize("cancelled", [False, True], ids=["publish-after-validation", "cancel-before-validation"])
def test_raw_is_persisted_while_display_and_message_remain_empty_until_validation(tmp_path, monkeypatch, cancelled):
    provider = ScriptedProviderDouble(pause=True)
    with typed_workspace(tmp_path, monkeypatch, provider) as workspace:
        query_id = post_query(workspace)
        try:
            assert provider.paused.wait(15), "The provider double must reach its explicit first-fragment barrier."
            events = stored_events(workspace, query_id)
            assert sum(event["type"] == "structured_delta" for event in events) == 1
            assert all(event["type"] not in {"delta", "done", "error", "cancelled"} for event in events)
            assert workspace.db.one("SELECT answer,state FROM query_runs WHERE id=?", (query_id,)) == {"answer": "", "state": "running"}
            assert workspace.db.rows("SELECT * FROM messages WHERE query_id=? AND role='assistant'", (query_id,)) == []
            detail = workspace.client.get(f"/api/v1/queries/{query_id}")
            assert detail.status_code == 200 and detail.json()["answer"] == ""
            assert "catalogue" not in detail.json()["metrics"] and "plan" not in detail.json()["metrics"]
            if cancelled:
                response = workspace.client.post(f"/api/v1/queries/{query_id}/cancel")
                assert response.status_code == 200
            else:
                workspace.client.portal.call(provider.unblock)
            await_query(workspace, query_id)
            if cancelled:
                stored = assert_no_unvalidated_answer(workspace, query_id, "cancelled")
                assert sum(event["type"] == "cancelled" for event in stored) == 1
                assert workspace.db.one("SELECT state FROM query_runs WHERE id=?", (query_id,))["state"] == "cancelled"
            else:
                stored = assert_raw_is_internal(workspace, query_id)
                assert stored[-1]["type"] == "done" and stored[-1]["data"]["text"]
        finally:
            if provider.gate is not None:
                workspace.client.portal.call(provider.unblock)
        assert workspace.app.state.queries.tasks == {} and workspace.app.state.queries.cancel_events == {}


def all_references(catalogue):
    return {"references": [reference["id"] for reference in catalogue["references"]], "comparisons": []}


def workbook_range(workspace):
    version_id = workspace.documents["xlsx"]["version_id"]
    unit = workspace.db.one("SELECT u.id,g.extraction_revision_id FROM office_units u "
                            "JOIN index_generations g ON g.id=u.generation_id WHERE g.version_id=?", (version_id,))
    return {"kind": "cell_range", "versionId": version_id, "extractionRevisionId": unit["extraction_revision_id"],
            "sheetId": unit["id"], "rowStart": 2, "rowEnd": 2, "columnStart": 2, "columnEnd": 3}


def test_cell_range_catalogue_and_reader_keep_exact_b2_c2_revision_without_other_cells(tmp_path, monkeypatch):
    provider = ScriptedProviderDouble(all_references)
    with typed_workspace(tmp_path, monkeypatch, provider) as workspace:
        scope = workbook_range(workspace)
        query_id = post_query(workspace, scope=scope)
        await_query(workspace, query_id)
        stored = assert_raw_is_internal(workspace, query_id)
        assert stored[-1]["type"] == "done" and stored[-1]["data"]["status"] == "done"
        answer = stored[-1]["data"]["text"]
        assert "72 V" in answer and "LEN(A2)" in answer and "OUTSIDE_SCOPE_999V" not in answer
        references = provider.catalogues[query_id]["references"]
        assert {reference["cell_facts"]["address"] for reference in references} == {"B2", "C2"}
        for reference in references:
            assert reference["kind"] == "cell"
            assert reference["version_id"] == scope["versionId"]
            assert reference["extraction_revision_id"] == scope["extractionRevisionId"]
            assert reference["locator"]["row_start"] == reference["locator"]["row_end"] == 2
            assert 2 <= reference["locator"]["column_start"] == reference["locator"]["column_end"] <= 3
        for source in stored[-1]["data"]["citations"]:
            assert {facts["address"] for facts in source["cell_facts"]} <= {"B2", "C2"}
            registered = workspace.client.get(f"/api/v1/citations/{query_id}/{source['source_id']}").json()
            assert registered["version_id"] == scope["versionId"] and registered["extraction_revision_id"] == scope["extractionRevisionId"]
        response = workspace.client.get(f"/api/v1/versions/{scope['versionId']}/sheets/{scope['sheetId']}/cells", params={
            "extraction_revision_id": scope["extractionRevisionId"], "row_start": 2, "row_end": 2, "column_start": 2, "column_end": 3})
        assert response.status_code == 200
        assert response.json()["extraction_revision_id"] == scope["extractionRevisionId"]
        cells = {cell["address"]: cell for cell in response.json()["cells"]}
        assert set(cells) == {"B2", "C2"} and cells["B2"]["display_text"] == "72 V"
        assert cells["C2"]["formula"]["effective_text"] == cells["C2"]["formula"]["raw_text"] == "LEN(A2)"
        assert cells["C2"]["cached_present"] is False
        assert cells["C2"]["cached_value"] is None


def test_render_budget_refuses_plan_without_publishing_or_prose_retry(tmp_path, monkeypatch):
    provider = ScriptedProviderDouble(all_references)
    with typed_workspace(tmp_path, monkeypatch, provider) as workspace:
        # Deliberate tiny test output budget; the product guard remains the actual renderer/token counter.
        workspace.app.state.settings.profile["llm"] = {"output_tokens_by_mode": {"ordinary": 1, "factual": 1}}
        query_id = post_query(workspace, scope=workbook_range(workspace))
        await_query(workspace, query_id)
        stored = assert_no_unvalidated_answer(workspace, query_id, "error")
        assert provider.calls[0]["output_tokens"] == 1
        assert stored[-1]["data"]["code"] == "structured_answer_render_budget_exceeded"


def test_legacy_pdf_docx_and_new_typed_history_replay_are_get_only_and_immutable(tmp_path, monkeypatch):
    provider = ScriptedProviderDouble()
    with typed_workspace(tmp_path, monkeypatch, provider) as workspace:
        legacy = []
        for document_format in ("pdf", "docx"):
            scope = {"kind": "documents", "documentIds": [workspace.documents[document_format]["document_id"]]}
            query_id = post_query(workspace, scope=scope)
            await_query(workspace, query_id)
            old_events = stored_events(workspace, query_id)
            assert all(event["type"] not in INTERNAL_EVENTS for event in old_events)
            assert old_events[-1]["data"]["text"] == "La tension nominale est 72 V [S001]."
            assert "structured_answer" not in old_events[-1]["data"]
            assert provider.calls[-1]["response_format"] is None
            legacy.append((query_id, scope, old_events))
        typed_id = post_query(workspace)
        await_query(workspace, typed_id)
        typed_events = assert_raw_is_internal(workspace, typed_id)
        assert typed_events[-1]["type"] == "done"
        # Archived citations remain tied to their original revision even with no active publication.
        workspace.db.execute("UPDATE documents SET active_generation_id=NULL")
        before = database_snapshot(workspace.db)
        calls_before = len(provider.calls)
        original_connect = workspace.db.connect

        @contextmanager
        def query_only():
            with original_connect() as connection:
                connection.execute("PRAGMA query_only=ON")
                yield connection

        def forbidden(*args, **kwargs):
            pytest.fail("History/citation GET must not resolve latest, retrieve, admit or create a query.")

        monkeypatch.setattr(workspace.db, "connect", query_only)
        monkeypatch.setattr(workspace.app.state.queries.resolver, "resolve", forbidden)
        monkeypatch.setattr(workspace.app.state.queries.search, "search", forbidden)
        monkeypatch.setattr(workspace.app.state.queries, "prepare_query", forbidden)
        monkeypatch.setattr(workspace.app.state.queries, "start_query", forbidden)
        listing = workspace.client.get("/api/v1/queries")
        assert listing.status_code == 200
        assert {query["query_id"] for query in listing.json()["queries"]} == {*[item[0] for item in legacy], typed_id}
        for query_id, expected_scope, expected_events in [*legacy, (typed_id, None, typed_events)]:
            response = workspace.client.get(f"/api/v1/queries/{query_id}")
            assert response.status_code == 200
            detail = response.json()
            assert detail["state"] == "done" and detail["mode"] is None
            assert detail["answer"] == expected_events[-1]["data"]["text"]
            assert "snapshot" not in detail and "catalogue" not in detail
            if expected_scope is not None:
                assert detail["scope"]["documentIds"] == expected_scope["documentIds"]
            replay = workspace.client.get(detail["events_url"], params={"after": 0})
            assert replay.status_code == 200
            assert sse_events(replay.text) == [event for event in expected_events if event["type"] not in INTERNAL_EVENTS]
            sources = next(event["data"]["sources"] for event in expected_events if event["type"] == "sources")
            for source in sources:
                citation = workspace.client.get(f"/api/v1/citations/{query_id}/{source['source_id']}")
                assert citation.status_code == 200
                for key in ("document_id", "version_id", "generation_id", "extraction_revision_id"):
                    assert citation.json()[key] == source[key]
                assert citation.json().get("locator") == source.get("locator")
        assert database_snapshot(workspace.db) == before
        assert len(provider.calls) == calls_before == 3
        assert workspace.app.state.queries.tasks == {} and workspace.app.state.queries.cancel_events == {}


def test_typed_literal_source_characters_and_citation_spans_remain_exact_in_delta_done_sql_and_history(tmp_path, monkeypatch):
    provider = ScriptedProviderDouble()
    with typed_workspace(tmp_path, monkeypatch, provider, literals=True) as workspace:
        query_id = post_query(workspace)
        await_query(workspace, query_id)
        stored = assert_raw_is_internal(workspace, query_id)
        assert stored[-1]["type"] == "done" and stored[-1]["data"]["status"] == "done"
        delta = next(event["data"] for event in stored if event["type"] == "delta")
        done = stored[-1]["data"]
        text = delta["text"]
        literal = "A😀 & <b> **[S999]** [S001,S002] ![x](https://example.invalid/x)"
        assert literal in text and "&amp;" not in text and "référence inconnue" not in text
        assert done["text"] == done["message"] == text
        assert workspace.db.one("SELECT answer FROM query_runs WHERE id=?", (query_id,))["answer"] == text
        assistant = workspace.db.one("SELECT content FROM messages WHERE query_id=? AND role='assistant'", (query_id,))
        assert assistant["content"] == text
        spans = delta["citation_spans"]
        assert spans == done["structured_answer"]["citation_spans"] and spans
        registered = {source["source_id"] for source in done["citations"]}
        cursor = 0
        for span in spans:
            start, end = span["start_offset"], span["end_offset"]
            assert cursor <= start < end <= len(text)
            assert text[start:end] == f"[{span['source_id']}]" and span["source_id"] in registered
            assert not (text.index(literal) <= start < text.index(literal) + len(literal))
            cursor = end
        detail = workspace.client.get(f"/api/v1/queries/{query_id}").json()
        assert detail["answer"] == text
        assert "citation_spans" not in detail["metrics"]["structured_answer"]


def test_snapshot_with_xlsx_keeps_typed_contract_when_retention_excludes_all_workbook_sources(tmp_path, monkeypatch):
    provider = ScriptedProviderDouble(all_references)
    with typed_workspace(tmp_path, monkeypatch, provider) as workspace:
        original_search = workspace.app.state.queries.search.search

        async def retained_source_filter_double(*args, **kwargs):
            # Explicit retention substitution: real search runs, then no workbook passage is supplied.
            result = await original_search(*args, **kwargs)
            result["results"] = [source for source in result["results"] if workspace.db.version(source["version_id"])["format"] != "xlsx"]
            result["warnings"].append({"code": "comparison_gap", "document_id": workspace.documents["xlsx"]["document_id"],
                                       "message": "Aucun passage XLSX retenu dans ce double de rétention."})
            return result

        monkeypatch.setattr(workspace.app.state.queries.search, "search", retained_source_filter_double)
        query_id = post_query(workspace, mode="comparison")
        await_query(workspace, query_id)
        stored = assert_raw_is_internal(workspace, query_id)
        done = stored[-1]["data"]
        assert stored[-1]["type"] == "done" and done["status"] == "done"
        assert len(provider.calls) == 1 and provider.calls[0]["response_format"] is not None
        references = provider.catalogues[query_id]["references"]
        assert references and all(reference["kind"] == "excerpt" for reference in references)
        assert {reference["document_id"] for reference in references} == {
            workspace.documents[format]["document_id"] for format in ("pdf", "docx")}
        assert all(workspace.db.version(source["version_id"])["format"] != "xlsx" for source in done["citations"])
        assert any(warning["code"] == "comparison_gap" for warning in done["warnings"])
        assert "LEN(A2)" not in done["text"]


def test_live_sse_skips_raw_internal_events_before_provider_validation(tmp_path, monkeypatch):
    provider = ScriptedProviderDouble(pause=True)
    with typed_workspace(tmp_path, monkeypatch, provider) as workspace:
        query_id = post_query(workspace)
        assert provider.paused.wait(15)
        assert workspace.db.one("SELECT state FROM query_runs WHERE id=?", (query_id,))["state"] == "running"
        original_rows = workspace.db.rows
        observed_stream = threading.Event()

        def observe_sse_read(sql, parameters=()):
            if sql.startswith("SELECT * FROM events ") and parameters and parameters[0] == query_id:
                observed_stream.set()
            return original_rows(sql, parameters)

        monkeypatch.setattr(workspace.db, "rows", observe_sse_read)

        async def release_after_live_read():
            for _ in range(100):
                if observed_stream.is_set():
                    await provider.unblock()
                    return
                await asyncio.sleep(.02)
            await provider.unblock()
            raise AssertionError("The actual SSE handler must read the running query before the provider is released.")

        future = workspace.client.portal.start_task_soon(release_after_live_read)
        try:
            response = workspace.client.get(f"/api/v1/queries/{query_id}/events", params={"after": 0})
            future.result(timeout=5)
            await_query(workspace, query_id)
            assert response.status_code == 200 and observed_stream.is_set()
            stored = stored_events(workspace, query_id)
            public = sse_events(response.text)
            assert public == [event for event in stored if event["type"] not in INTERNAL_EVENTS]
            assert public[-1]["type"] == "done" and public[-1]["data"]["status"] == "done"
            assert "event: structured_" not in response.text
            assert "".join(event["data"]["text"] for event in stored if event["type"] == "structured_delta") == provider.raw_by_query[query_id]

            def strings(value):
                if isinstance(value, str):
                    yield value
                elif isinstance(value, dict):
                    for item in value.values():
                        yield from strings(item)
                elif isinstance(value, list):
                    for item in value:
                        yield from strings(item)

            assert all(provider.raw_by_query[query_id] not in value for value in strings(public))
        finally:
            if provider.gate is not None:
                workspace.client.portal.call(provider.unblock)
            future.result(timeout=5)
