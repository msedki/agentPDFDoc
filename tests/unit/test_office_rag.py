"""Real OOXML/SQLite/FTS; embeddings and vector transport are explicit doubles."""

import asyncio
import hashlib
import json
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook
from test_api_storage import FakeEmbedding, FakeLlmTokenizer, FakeVectors

from services.api.context import ContextBuilder
from services.api.db import Database, now, uid
from services.api.errors import ApiError
from services.api.indexing import Indexer
from services.api.office_search import projection_fragments
from services.api.query import QueryService
from services.api.retrieval import SearchService
from services.api.schemas import QueryRequest, Scope
from services.api.scope import ScopeResolver
from services.api.settings import Settings
from services.ingestion.office.pipeline import extract_office


class RecordingEmbedding(FakeEmbedding):
    def __init__(self):
        self.submitted = []

    def embed(self, texts, passage=True):
        self.submitted.extend(texts)
        return super().embed(texts, passage)


@pytest.fixture
def rag(tmp_path):
    settings = Settings(tmp_path, {"app": {"data_dir": "runtime", "port": 8785}})
    db = Database(settings.db_path)
    db.initialize()
    embedding, vectors = RecordingEmbedding(), FakeVectors()
    indexer = Indexer(db, embedding, vectors, settings, FakeLlmTokenizer())
    resolver = ScopeResolver(db)
    search = SearchService(db, resolver, embedding, vectors, settings, projector=indexer)
    return settings, db, embedding, indexer, resolver, search


def workbook(rag, name="budget", outside="outside-secret", long=False):
    settings, db, _, indexer, _, _ = rag
    path = settings.data_dir / "originals" / (name + ".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Budget été"
    sheet.append(["Allowed header", "Excluded header"])
    sheet.append(["CCU-21 : tension 72 V. A😀é e\u0301", outside])
    sheet.append(["CCU-22 : tension 48 V.", outside])
    sheet.append([("CCU-23 : tension 24 V. " * 500) if long else "CCU-23 : tension 24 V.", outside])
    book.create_sheet("Autre").append(["Another-sheet-secret"])
    book.save(path)
    book.close()
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    imported = db.import_original(name + ".xlsx", digest, path)
    extraction = extract_office(path, settings.data_dir / "extractions" / imported["version_id"], {}, imported["version_id"], "xlsx")
    asyncio.run(indexer.index(imported["job_id"], extraction))
    return imported, extraction


def range_scope(imported, extraction, **changes):
    values = {"kind": "cell_range", "versionId": imported["version_id"], "extractionRevisionId": extraction["extraction_revision_id"],
              "sheetId": extraction["units"][0]["id"], "rowStart": 2, "rowEnd": 4, "columnStart": 1, "columnEnd": 1}
    return Scope(**{**values, **changes})


def test_range_projection_precedes_lexical_dense_and_context(rag):
    imported, extraction = workbook(rag, outside="excluded-secret CCU-99 999 V " * 500)
    _, _, embedding, _, resolver, search = rag
    snapshot = resolver.resolve(range_scope(imported, extraction))
    embedding.submitted.clear()
    result = asyncio.run(search.search("Quelle tension CCU-21 ?", snapshot))
    assert result["results"]
    assert result["results"][0]["locator"]["cell_range"] == "A2"
    assert not any("excluded-secret" in text or "Excluded header" in text for text in embedding.submitted)
    assert embedding.submitted[0] == "Quelle tension CCU-21 ?"
    assert "CCU-99" not in json.dumps(result)
    for source in result["results"]:
        assert source["page_indices"] == [] and source["page_index"] is None
        assert source["scope_projected"] and source["format"] == "xlsx"
        evidence = json.loads(ContextBuilder.evidence({**source, "source_id": "S001"}))
        assert evidence["pages"] == [] and evidence["locators"][0]["column_start"] == 1
        assert "cache" in evidence["limitations"]
        expanded = resolver.expand_parent(source, snapshot, FakeLlmTokenizer())
        assert expanded == source
        for block in source["blocks"]:
            raw = rag[1].one("SELECT text,source_text_hash FROM blocks WHERE generation_id=? AND id=?", (source["generation_id"], block["id"]))
            assert raw["text"][block["start_offset"]:block["end_offset"]] == block["text"]
            assert raw["source_text_hash"] == block["source_text_hash"]


def test_outside_only_changes_cannot_change_authorized_ranks_or_identifiers(rag):
    first, first_extraction = workbook(rag, "first", "small outside")
    second, second_extraction = workbook(rag, "second", "CCU-21 tension CCU-99 SECRET " * 1000)
    resolver, search = rag[4:]
    answers = [asyncio.run(search.search("CCU-21 tension", resolver.resolve(range_scope(imported, extraction))))
               for imported, extraction in ((first, first_extraction), (second, second_extraction))]
    def oracle(answer):
        return [(source["locator"]["cell_range"], source["text"], source["score"], source["exact_identifier"]) for source in answer["top10"]]
    assert oracle(answers[0]) == oracle(answers[1])
    assert not any(warning.get("identifier") == "CCU-99" for answer in answers for warning in answer["warnings"])


def test_projection_chunks_long_cells_with_exact_source_spans(rag):
    imported, extraction = workbook(rag, long=True, outside="other" * 5000)
    snapshot = rag[4].resolve(range_scope(imported, extraction, rowStart=4, rowEnd=4))
    sources = projection_fragments(rag[1], snapshot, rag[2])
    assert len(sources) > 2
    assert all(rag[2].count(source["text"]) <= 320 for source in sources)
    assert all(source["locator"]["cell_range"] == "A4" for source in sources)
    assert not any("other" in source["text"] for source in sources)
    for source in sources:
        for block in source["blocks"]:
            raw = rag[1].one("SELECT text FROM blocks WHERE generation_id=? AND id=?", (source["generation_id"], block["id"]))["text"]
            assert block["text"] == raw[block["start_offset"]:block["end_offset"]]


def test_range_uses_cached_vectors_with_the_full_existing_identity(rag):
    imported, extraction = workbook(rag)
    snapshot = rag[4].resolve(range_scope(imported, extraction))
    for _ in range(2):
        rag[2].submitted.clear()
        asyncio.run(rag[5].search("CCU-21", snapshot))
        submitted = list(rag[2].submitted)
    assert submitted == ["CCU-21"]
    assert rag[1].one("SELECT count(*) n FROM embedding_cache")["n"] > 0


def test_unknown_sheet_or_revision_is_never_substituted(rag):
    imported, extraction = workbook(rag)
    for changes, code in (({"sheetId": "not-a-sheet"}, "sheet_not_found"), ({"extractionRevisionId": uid()}, "extraction_revision_not_found")):
        with pytest.raises(ApiError) as caught:
            rag[4].resolve(range_scope(imported, extraction, **changes))
        assert caught.value.code == code


def test_empty_cell_range_returns_empty_evidence(rag):
    imported, extraction = workbook(rag)
    snapshot = rag[4].resolve(range_scope(imported, extraction, rowStart=100, rowEnd=101))
    result = asyncio.run(rag[5].search("CCU-21", snapshot))
    assert result["results"] == []


def test_direct_chunk_lookup_cannot_bypass_required_range_projection(rag):
    imported, extraction = workbook(rag)
    snapshot = rag[4].resolve(range_scope(imported, extraction))
    chunk = rag[1].one("SELECT * FROM chunks LIMIT 1")
    with pytest.raises(ApiError, match="projection"):
        rag[4].source_for_chunk(chunk, snapshot)


def test_sheet_scope_preserves_native_locations_and_other_sheet_exclusion(rag):
    imported, extraction = workbook(rag)
    scope = Scope(kind="sheet", versionId=imported["version_id"], extractionRevisionId=extraction["extraction_revision_id"], sheetId=extraction["units"][0]["id"])
    snapshot = rag[4].resolve(scope)
    sources = [rag[4].source_for_chunk(chunk, snapshot) for chunk in rag[1].rows("SELECT * FROM chunks")]
    sources = [source for source in sources if source]
    assert sources and all(source["locator"]["sheet_id"] == scope.sheetId for source in sources)
    assert "Another-sheet-secret" not in json.dumps(sources)


def test_xlsx_selection_clips_cell_bindings_and_citation_to_authorized_offsets(rag):
    imported, extraction = workbook(rag, outside="excluded neighbour")
    block = next(block for block in extraction["units"][0]["blocks"]
                 if any(binding["address"] == "A2" for binding in block["bindings"]))
    binding = next(binding for binding in block["bindings"] if binding["address"] == "A2")
    start, end = binding["start"] + 1, binding["end"] - 1
    scope = Scope(kind="selection", versionId=imported["version_id"], spans=[{
        "blockId": block["id"], "extractionRevisionId": extraction["extraction_revision_id"],
        "blockTextSha256": block["source_text_hash"], "offsetUnit": "unicode_code_point",
        "startOffset": start, "endOffset": end,
    }])
    snapshot = rag[4].resolve(scope)
    sources = rag[4].selected_sources(snapshot)
    sources.extend(source for chunk in rag[1].rows("SELECT * FROM chunks")
                   if (source := rag[4].source_for_chunk(chunk, snapshot)))
    assert sources
    for source in sources:
        assert source["locator"]["cell_range"] == "A2"
        assert source["precision"] == "cell"
        assert [fact["address"] for fact in source["cell_facts"]] == ["A2"]
        assert source["cell_facts"][0]["value_origin"] == "literal"
        for selected in source["blocks"]:
            assert selected["structure"]["addresses"] == ["A2"]
            assert selected["structure"]["table_contexts"] == []
            assert {item["address"] for item in selected["bindings"]} == {"A2"}
            for item in selected["bindings"]:
                assert selected["start_offset"] <= item["start"] < item["end"] <= selected["end_offset"]
                assert item["source_start"] == binding["source_start"] + item["start"] - binding["start"]
                assert item["source_end"] - item["source_start"] == item["end"] - item["start"]
        assert "excluded neighbour" not in json.dumps(source)


def test_registered_cell_citation_has_exact_revision_and_no_pdf_page(rag):
    imported, extraction = workbook(rag)
    snapshot = rag[4].resolve(range_scope(imported, extraction))
    source = asyncio.run(rag[5].search("CCU-21", snapshot))["results"][0]
    query_id = uid()
    rag[1].execute("INSERT INTO query_runs(id,question,scope_json,snapshot_json,state,created_at,updated_at) VALUES(?,?,?,?,'complete',?,?)", (query_id, "CCU-21", "{}", "{}", now(), now()))
    registered = QueryService.register_sources(SimpleNamespace(db=rag[1]), query_id, [{**source, "source_id": "S001"}])[0]
    assert registered["page_index"] is None and registered["page_number"] is None
    assert registered["precision"] == "cell" and registered["locator"]["cell_range"] == "A2"
    assert registered["extraction_revision_id"] == extraction["extraction_revision_id"]
    assert json.loads(rag[1].one("SELECT source_json FROM citations")["source_json"]) == registered


def test_docx_parent_section_includes_nested_section_source_blocks(rag):
    settings, db, _, indexer, resolver, _ = rag
    path = settings.data_dir / "originals" / "hierarchy.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("Parent", 1)
    doc.add_paragraph("Parent proof")
    doc.add_heading("Child", 2)
    doc.add_paragraph("Child proof CCU-99")
    doc.add_heading("Other", 1)
    doc.add_paragraph("Other secret")
    doc.save(path)
    imported = db.import_original("hierarchy.docx", hashlib.sha256(path.read_bytes()).hexdigest(), path)
    extraction = extract_office(path, settings.data_dir / "extractions" / imported["version_id"], {}, imported["version_id"], "docx")
    asyncio.run(indexer.index(imported["job_id"], extraction))
    parent = next(section for section in extraction["sections"] if section["title"] == "Parent")
    assert len(parent["block_ids"]) == 4
    snapshot = resolver.resolve(Scope(kind="section", versionId=imported["version_id"], extractionRevisionId=extraction["extraction_revision_id"], sectionId=parent["id"]))
    assert snapshot.block_ids == parent["block_ids"]
    sources = [resolver.source_for_chunk(chunk, snapshot) for chunk in db.rows("SELECT * FROM chunks")]
    text = "\n".join(source["text"] for source in sources if source)
    assert "Child proof CCU-99" in text and "Other secret" not in text


def test_cell_citation_focus_keeps_exact_cell_within_a_larger_allowed_range(rag):
    imported, extraction = workbook(rag, outside="neighbour CCU-99")
    original = rag[4].resolve(range_scope(imported, extraction, rowStart=2, rowEnd=2))
    source = asyncio.run(rag[5].search("CCU-21", original))["results"][0]
    query_id = uid()
    rag[1].execute("INSERT INTO query_runs(id,question,scope_json,snapshot_json,state,created_at,updated_at) VALUES(?,?,?,?,'complete',?,?)", (query_id, "CCU-21", "{}", "{}", now(), now()))
    QueryService.register_sources(SimpleNamespace(db=rag[1]), query_id, [{**source, "source_id": "S001"}])
    scope = range_scope(imported, extraction, rowStart=2, rowEnd=2, columnEnd=2)
    snapshot = rag[4].resolve(scope)
    request = QueryRequest(question="Précise la tension", scope=scope, focus={"query_id": query_id, "source_id": "S001"})
    QueryService.resolve_followup(SimpleNamespace(db=rag[1]), request, snapshot)
    fragments = projection_fragments(rag[1], snapshot, rag[2])
    assert fragments and {source["locator"]["cell_range"] for source in fragments} == {"A2"}
    assert "neighbour" not in json.dumps(fragments)
    assert snapshot.as_dict()["cell_ranges"] == [{"row_start": 2, "row_end": 2, "column_start": 1, "column_end": 1}]


def test_table_headers_reach_full_sheet_context_but_are_excluded_from_cell_projection(rag):
    from test_office_indexing import ingest_fixture
    storage = rag[0], rag[1], FakeVectors(), rag[3]
    imported, extraction = ingest_fixture(storage, "xlsx")
    asyncio.run(rag[3].index(imported["job_id"], extraction))
    sheet = Scope(kind="sheet", versionId=imported["version_id"], extractionRevisionId=extraction["extraction_revision_id"], sheetId=extraction["units"][0]["id"])
    snapshot = rag[4].resolve(sheet)
    source = next(source for chunk in rag[1].rows("SELECT * FROM chunks")
                  if (source := rag[4].source_for_chunk(chunk, snapshot)) and "9007199254740993" in source["text"])
    evidence = json.loads(ContextBuilder.evidence({**source, "source_id": "S001"}))
    assert "Amount" in json.dumps(evidence["table_contexts"])
    assert evidence["cell_facts"]
    narrow = rag[4].resolve(range_scope(imported, extraction, rowStart=2, rowEnd=2, columnStart=2, columnEnd=2))
    projected = projection_fragments(rag[1], narrow, rag[2])[0]
    limited = json.loads(ContextBuilder.evidence({**projected, "source_id": "S001"}))
    assert "table_contexts" not in limited and "Amount" not in json.dumps(limited)
    assert limited["cell_facts"][0]["address"] == "B2"


def test_context_distinguishes_literal_cells_from_formula_caches_and_missing_results(rag):
    from test_office_indexing import ingest_fixture
    imported, extraction = ingest_fixture((rag[0], rag[1], FakeVectors(), rag[3]), "xlsx")
    asyncio.run(rag[3].index(imported["job_id"], extraction))
    for address, row, column, origin in (("B2", 2, 2, "literal"), ("C3", 3, 3, "formula_without_cache"),
                                         ("C4", 4, 3, "formula_cache")):
        snapshot = rag[4].resolve(range_scope(imported, extraction, rowStart=row, rowEnd=row,
                                             columnStart=column, columnEnd=column))
        source = projection_fragments(rag[1], snapshot, rag[2])[0]
        evidence = json.loads(ContextBuilder.evidence({**source, "source_id": "S001"}))
        facts = evidence["cell_facts"][0]
        assert facts["address"] == address
        assert facts["value_origin"] == origin
        if origin == "literal":
            assert not facts["formula_present"] and "littérales" in evidence["limitations"]
            assert "valeurs calculées uniquement" not in evidence["limitations"]
        else:
            assert facts["formula_present"] and "jamais exécutées" in evidence["limitations"]
            assert facts["cached_present"] is (origin == "formula_cache")


def test_empty_commented_cell_is_not_described_as_a_literal_value(rag):
    from openpyxl.comments import Comment

    from services.api.office import cell_semantics
    path = rag[0].data_dir / "originals" / "empty-comment.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Workbook()
    document.active["B1"].comment = Comment("Cellule vide documentée", "QA")
    document.save(path)
    imported = rag[1].import_original(path.name, hashlib.sha256(path.read_bytes()).hexdigest(), path)
    extraction = extract_office(path, rag[0].data_dir / "extractions" / imported["version_id"], {}, imported["version_id"], "xlsx")
    asyncio.run(rag[3].index(imported["job_id"], extraction))
    empty = rag[1].one("SELECT data_json FROM office_cells WHERE address='B1'")
    facts = cell_semantics(json.loads(empty["data_json"]))
    assert facts["value_origin"] == "empty" and facts["value_kind"] == "empty" and not facts["formula_present"]
